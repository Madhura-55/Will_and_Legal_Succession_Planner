from docx import Document
from will_schema import WillData
from datetime import date

def build_will(data: WillData) -> Document:
    doc = Document()
    doc.add_heading("Last Will and Testament", 0)
    doc.add_paragraph(
        f"I, {data.testator_name or '[Name]'}, aged {data.testator_age or '[Age]'} years, "
        f"residing at {data.testator_address or '[Address]'}, being of sound mind and memory, "
        f"hereby revoke all former wills and codicils made by me and declare this to be my "
        f"Last Will and Testament, made this {date.today().strftime('%d %B %Y')}. "
        f"This document was prepared with the assistance of Smart-Will."
    )

    doc.add_heading("Assets", level=1)
    for a in data.assets:
        doc.add_paragraph(
            f"• {a.asset_type.title()}: {a.description}" +
            (f" ({a.identifier})" if a.identifier else ""),
            style="List Bullet"
        )

    doc.add_heading("Beneficiaries", level=1)
    for b in data.beneficiaries:
        doc.add_paragraph(
            f"I bequeath {b.allocation} to {b.name} ({b.relationship})." +
            (f" As {b.name} is a minor, {data.minor_guardian or '[Guardian]'} shall act as guardian."
             if b.is_minor else ""),
            style="List Bullet"
        )

    doc.add_heading("Executor", level=1)
    doc.add_paragraph(
        f"I appoint {data.executor_name or '[Executor]'} "
        f"({data.executor_relationship or ''}) as the Executor of this Will."
    )

    doc.add_heading("Attestation", level=1)
    doc.add_paragraph(
        "Signed and declared by the above-named Testator as their Last Will, "
        "in the presence of us, who at their request, in their presence, "
        "and in the presence of each other, have subscribed our names as witnesses."
    )
    for i, w in enumerate(data.witnesses, 1):
        doc.add_paragraph(f"Witness {i}: {w}")

    doc.add_paragraph("\n\nTestator signature: _____________________")
    return doc
